from __future__ import annotations

import hashlib
import json
import zipfile
from collections.abc import Iterable
from io import BytesIO
from pathlib import Path
from typing import Any

from inventory_optimizer.domain.requests import OptimizationRequest
from openpyxl import load_workbook
from pydantic import ValidationError

from .contracts import DatasetEnvelope, LocatedIssue, SourceDisclosure
from .datasets import envelope_from_request
from .field_registry import ROW_LIMITS, SHEETS, FieldSpec

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
ALLOWED_SHEETS = {"README", "data_dictionary", *SHEETS}
FORBIDDEN_ARCHIVE_PARTS = ("vbaproject.bin", "externallinks/", "embeddings/", "oleobject", "connections.xml")


class UploadValidationError(ValueError):
    def __init__(self, issues: Iterable[LocatedIssue]):
        self.issues = tuple(issues)
        super().__init__(f"upload rejected with {len(self.issues)} located issue(s)")


def _duplicate_rejecting_pairs(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in pairs:
        if key in result:
            raise UploadValidationError(
                [LocatedIssue(code="DUPLICATE_JSON_KEY", message="Duplicate JSON key.", location=f"$.{key}")]
            )
        result[key] = value
    return result


def _pydantic_issues(exc: ValidationError, prefix: str = "$") -> tuple[LocatedIssue, ...]:
    issues: list[LocatedIssue] = []
    for error in exc.errors(include_url=False, include_context=False, include_input=False):
        path = prefix
        for part in error["loc"]:
            path += f"[{part}]" if isinstance(part, int) else f".{part}"
        issues.append(
            LocatedIssue(
                code=f"CONTRACT_{str(error['type']).upper()}",
                message=str(error["msg"]),
                location=path,
            )
        )
    return tuple(issues)


def _unsupported_feature_issues(request: OptimizationRequest) -> tuple[LocatedIssue, ...]:
    issues: list[LocatedIssue] = []
    if request.planning_periods:
        issues.append(LocatedIssue(code="CAPABILITY_MULTI_PERIOD_DISABLED", message="Multi-period input is not enabled.", location="$.planning_periods"))
    if request.known_future_events:
        issues.append(LocatedIssue(code="CAPABILITY_EVENTS_DISABLED", message="Known future events are not enabled.", location="$.known_future_events"))
    if request.expected_economics:
        issues.append(LocatedIssue(code="CAPABILITY_EXPECTED_ECONOMICS_DISABLED", message="Expected economics is not enabled.", location="$.expected_economics"))
    if any(row.candidate_fee_rates for row in request.demand):
        issues.append(LocatedIssue(code="CAPABILITY_FEE_TIERS_DISABLED", message="Fee-tier authoring is not enabled.", location="$.demand"))
    return tuple(issues)


def _build_envelope(request: OptimizationRequest, *, data: bytes, filename: str) -> DatasetEnvelope:
    limit_issues: list[LocatedIssue] = []
    for section, limit in ROW_LIMITS.items():
        if section == "request":
            continue
        rows = getattr(request, section)
        if len(rows) > limit:
            limit_issues.append(
                LocatedIssue(
                    code="ROW_LIMIT",
                    message=f"{section} exceeds the {limit} row limit.",
                    location=f"$.{section}",
                )
            )
    if limit_issues:
        raise UploadValidationError(limit_issues)
    unsupported = _unsupported_feature_issues(request)
    if unsupported:
        raise UploadValidationError(unsupported)
    digest = hashlib.sha256(data).hexdigest()
    return envelope_from_request(
        request,
        source=SourceDisclosure(
            mode="upload",
            label="Uploaded",
            upload_hash=digest,
            filename=Path(filename).name,
        ),
    )


def parse_json_upload(data: bytes, filename: str) -> DatasetEnvelope:
    try:
        payload = json.loads(data.decode("utf-8-sig"), object_pairs_hook=_duplicate_rejecting_pairs)
    except UnicodeDecodeError as exc:
        raise UploadValidationError(
            [LocatedIssue(code="JSON_ENCODING", message="JSON must be UTF-8.", location="$")]
        ) from exc
    except json.JSONDecodeError as exc:
        raise UploadValidationError(
            [LocatedIssue(code="JSON_SYNTAX", message=exc.msg, location=f"line {exc.lineno}, column {exc.colno}")]
        ) from exc
    if not isinstance(payload, dict):
        raise UploadValidationError(
            [LocatedIssue(code="JSON_ROOT_TYPE", message="JSON root must be an object.", location="$")]
        )
    try:
        request = OptimizationRequest.model_validate(payload)
    except ValidationError as exc:
        raise UploadValidationError(_pydantic_issues(exc)) from exc
    return _build_envelope(request, data=data, filename=filename)


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _convert_cell(value: Any, spec: FieldSpec, location: str) -> tuple[Any, LocatedIssue | None]:
    if _is_blank(value):
        if spec.required:
            return None, LocatedIssue(code="REQUIRED_CELL_EMPTY", message="Required value is empty.", location=location)
        return None, None
    try:
        if spec.data_type == "string":
            if not isinstance(value, str):
                raise TypeError("expected text")
            return value.strip(), None
        if spec.data_type in {"datetime", "date"}:
            return value.isoformat() if hasattr(value, "isoformat") else str(value).strip(), None
        if spec.data_type == "boolean":
            if isinstance(value, bool):
                return value, None
            if isinstance(value, str) and value.strip().lower() in {"true", "false"}:
                return value.strip().lower() == "true", None
            raise TypeError("expected true or false")
        if spec.data_type == "integer":
            if isinstance(value, bool) or int(value) != float(value):
                raise TypeError("expected an integer")
            return int(value), None
        if spec.data_type == "number":
            if isinstance(value, bool):
                raise TypeError("expected a number")
            return float(value), None
        if spec.data_type == "json":
            return json.loads(value) if isinstance(value, str) else value, None
        return value, None
    except (TypeError, ValueError, json.JSONDecodeError) as exc:
        return None, LocatedIssue(code="CELL_TYPE", message=str(exc), location=location)


def _check_archive(data: bytes) -> tuple[LocatedIssue, ...]:
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            lowered = [name.lower() for name in archive.namelist()]
    except zipfile.BadZipFile:
        return (LocatedIssue(code="XLSX_ARCHIVE_INVALID", message="File is not a valid XLSX archive.", location="$file"),)
    return tuple(
        LocatedIssue(code="XLSX_UNSAFE_CONTENT", message="Workbook contains unsupported active or linked content.", location=f"$file:{name}")
        for name in lowered
        if any(part in name for part in FORBIDDEN_ARCHIVE_PARTS)
    )


def _sheet_rows(sheet: Any, sheet_name: str, specs: tuple[FieldSpec, ...]) -> tuple[list[dict[str, Any]], list[LocatedIssue]]:
    issues: list[LocatedIssue] = []
    rows: list[dict[str, Any]] = []
    raw_headers = [cell.value for cell in sheet[1]]
    while raw_headers and _is_blank(raw_headers[-1]):
        raw_headers.pop()
    headers = [str(value).strip() if value is not None else "" for value in raw_headers]
    expected = [spec.name for spec in specs]
    if len(headers) != len(set(headers)):
        issues.append(LocatedIssue(code="DUPLICATE_COLUMN", message="Duplicate column names are not allowed.", location=f"{sheet_name}!1"))
    unknown = [value for value in headers if value and value not in expected]
    missing = [value for value in expected if value not in headers]
    for value in unknown:
        issues.append(LocatedIssue(code="UNKNOWN_COLUMN", message="Unknown columns are not allowed.", location=f"{sheet_name}!{value}"))
    for value in missing:
        issues.append(LocatedIssue(code="MISSING_COLUMN", message="Template column is missing.", location=f"{sheet_name}!{value}"))
    if unknown or missing or len(headers) != len(set(headers)):
        return rows, issues
    specs_by_name = {spec.name: spec for spec in specs}
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=len(headers)), start=2):
        if all(_is_blank(cell.value) for cell in cells):
            continue
        record: dict[str, Any] = {}
        for column_number, cell in enumerate(cells, start=1):
            name = headers[column_number - 1]
            if cell.data_type == "f":
                issues.append(LocatedIssue(code="FORMULA_NOT_ALLOWED", message="Formulas are not allowed.", location=f"{sheet_name}!{cell.coordinate}"))
                continue
            converted, issue = _convert_cell(cell.value, specs_by_name[name], f"{sheet_name}!{cell.coordinate}")
            if issue:
                issues.append(issue)
            elif converted is not None:
                record[name] = converted
        rows.append(record)
    limit = ROW_LIMITS.get(sheet_name)
    if limit is not None and len(rows) > limit:
        issues.append(LocatedIssue(code="ROW_LIMIT", message=f"Sheet exceeds the {limit} row limit.", location=sheet_name))
    return rows, issues


def _request_sheet(sheet: Any) -> tuple[dict[str, Any], list[LocatedIssue]]:
    issues: list[LocatedIssue] = []
    headers = [cell.value for cell in sheet[1]]
    if headers[:2] != ["field", "value"]:
        return {}, [LocatedIssue(code="REQUEST_HEADERS", message="Request sheet must begin with field,value.", location="request!1")]
    allowed = {spec.name: spec for spec in SHEETS["request"]}
    values: dict[str, Any] = {}
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=2), start=2):
        key = cells[0].value
        value = cells[1].value
        if _is_blank(key) and _is_blank(value):
            continue
        key = str(key).strip()
        if key not in allowed:
            issues.append(LocatedIssue(code="UNKNOWN_REQUEST_FIELD", message="Unknown request field.", location=f"request!A{row_number}"))
            continue
        if key in values:
            issues.append(LocatedIssue(code="DUPLICATE_REQUEST_FIELD", message="Duplicate request field.", location=f"request!A{row_number}"))
            continue
        if cells[1].data_type == "f":
            issues.append(LocatedIssue(code="FORMULA_NOT_ALLOWED", message="Formulas are not allowed.", location=f"request!B{row_number}"))
            continue
        converted, issue = _convert_cell(value, allowed[key], f"request!B{row_number}")
        if issue:
            issues.append(issue)
        else:
            values[key] = converted
    for key, spec in allowed.items():
        if spec.required and key not in values:
            issues.append(LocatedIssue(code="MISSING_REQUEST_FIELD", message="Required request field is missing.", location=f"request!{key}"))
    if values.get("template_schema_version") != "1.0":
        issues.append(LocatedIssue(code="TEMPLATE_VERSION", message="Only template schema 1.0 is supported.", location="request!template_schema_version"))
    return values, issues


def parse_xlsx_upload(data: bytes, filename: str) -> DatasetEnvelope:
    archive_issues = _check_archive(data)
    if archive_issues:
        raise UploadValidationError(archive_issues)
    try:
        workbook = load_workbook(BytesIO(data), read_only=False, data_only=False, keep_links=False)
    except Exception as exc:
        raise UploadValidationError(
            [LocatedIssue(code="XLSX_OPEN", message="Workbook could not be opened.", location="$file")]
        ) from exc
    issues: list[LocatedIssue] = []
    unknown_sheets = sorted(set(workbook.sheetnames) - ALLOWED_SHEETS)
    for name in unknown_sheets:
        issues.append(LocatedIssue(code="UNKNOWN_SHEET", message="Unknown sheets are not allowed.", location=name))
    for required_sheet in SHEETS:
        if required_sheet not in workbook.sheetnames:
            issues.append(LocatedIssue(code="MISSING_SHEET", message="Required template sheet is missing.", location=required_sheet))
    if issues:
        raise UploadValidationError(issues)

    request_values, request_issues = _request_sheet(workbook["request"])
    issues.extend(request_issues)
    payload: dict[str, Any] = {
        key: value for key, value in request_values.items() if key != "template_schema_version"
    }
    for sheet_name, specs in SHEETS.items():
        if sheet_name == "request":
            continue
        rows, row_issues = _sheet_rows(workbook[sheet_name], sheet_name, specs)
        issues.extend(row_issues)
        if sheet_name == "entity_relationships":
            rows = [
                {
                    "value": {
                        key: row.pop(key)
                        for key in (
                            "entity_id",
                            "legal_entity_id",
                            "ultimate_parent_id",
                            "relationship_type",
                            "ownership_confidence",
                            "approval_reference",
                        )
                        if key in row
                    },
                    **row,
                }
                for row in rows
            ]
        payload[sheet_name] = rows
    if issues:
        raise UploadValidationError(issues)
    try:
        request = OptimizationRequest.model_validate(payload)
    except ValidationError as exc:
        raise UploadValidationError(_pydantic_issues(exc, "$workbook")) from exc
    return _build_envelope(request, data=data, filename=filename)


def parse_upload(data: bytes, filename: str) -> DatasetEnvelope:
    if not data:
        raise UploadValidationError(
            [LocatedIssue(code="EMPTY_UPLOAD", message="Upload is empty.", location="$file")]
        )
    if len(data) > MAX_UPLOAD_BYTES:
        raise UploadValidationError(
            [LocatedIssue(code="UPLOAD_SIZE", message="Upload exceeds 10 MiB.", location="$file")]
        )
    suffix = Path(filename).suffix.lower()
    if suffix == ".json":
        return parse_json_upload(data, filename)
    if suffix == ".xlsx":
        return parse_xlsx_upload(data, filename)
    raise UploadValidationError(
        [LocatedIssue(code="UPLOAD_TYPE", message="Only .json and .xlsx are supported.", location="$file")]
    )
