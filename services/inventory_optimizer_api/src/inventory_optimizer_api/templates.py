from __future__ import annotations

import json
from io import BytesIO
from typing import Any, Literal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .field_registry import SHEETS
from .synthetic import generate_request, starter_config


def _json_cell(value: Any) -> Any:
    if isinstance(value, (list, dict)):
        return json.dumps(value, sort_keys=True)
    return value


def _entity_row(row: dict[str, Any]) -> dict[str, Any]:
    return {**row["value"], **{key: value for key, value in row.items() if key != "value"}}


def _request_rows(payload: dict[str, Any]) -> list[tuple[str, Any]]:
    return [
        ("template_schema_version", "1.0"),
        ("request_id", payload["request_id"]),
        ("as_of", payload["as_of"]),
        ("effective_date", payload["effective_date"]),
        ("problem_family", payload["problem_family"]),
    ]


def workbook_template(kind: Literal["starter-xlsx", "advanced-xlsx"]) -> bytes:
    request, _ = generate_request(starter_config())
    payload = request.model_dump(mode="json")
    include_examples = kind == "starter-xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)

    readme = workbook.create_sheet("README")
    readme.append(["InventoryBallast IBAL input template"])
    readme.append(["Schema", "1.0"])
    readme.append(["Format", "Exact units; formulas, macros, external links and unknown sheets/columns are rejected."])
    readme.append(["Balance identity", "available_to_lend_shares = total_lendable_shares - on_loan_shares - reserved_shares - committed_out_shares"])
    readme.append(["Rates", "Decimal annual rates: 0.02 means 2%."])
    readme["A1"].font = Font(bold=True, color="FFFFFF")
    readme["A1"].fill = PatternFill("solid", fgColor="1D4ED8")
    readme.column_dimensions["A"].width = 24
    readme.column_dimensions["B"].width = 110

    dictionary = workbook.create_sheet("data_dictionary")
    dictionary.append(["sheet", "field", "type", "required", "unit", "description", "example"])
    for sheet_name, fields in SHEETS.items():
        for spec in fields:
            dictionary.append([sheet_name, spec.name, spec.data_type, spec.required, spec.unit, spec.description, _json_cell(spec.example)])
    for cell in dictionary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
    dictionary.freeze_panes = "A2"
    dictionary.auto_filter.ref = dictionary.dimensions

    for sheet_name, fields in SHEETS.items():
        sheet = workbook.create_sheet(sheet_name)
        headers = [spec.name for spec in fields]
        if sheet_name == "request":
            sheet.append(["field", "value"])
            if include_examples:
                for key, value in _request_rows(payload):
                    sheet.append([key, _json_cell(value)])
        else:
            sheet.append(headers)
            if include_examples:
                rows = payload.get(sheet_name, [])
                if sheet_name == "entity_relationships":
                    rows = [_entity_row(row) for row in rows]
                for row in rows:
                    sheet.append([_json_cell(row.get(header)) for header in headers])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F766E")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def json_template() -> bytes:
    request, _ = generate_request(starter_config())
    return request.model_dump_json(indent=2).encode("utf-8")
