from __future__ import annotations

import json
import zipfile
from io import BytesIO

import pytest
from openpyxl import load_workbook

from inventory_optimizer_api.templates import json_template, workbook_template
from inventory_optimizer_api.upload import UploadValidationError, parse_upload


def _issue_codes(exc_info: pytest.ExceptionInfo[UploadValidationError]) -> set[str]:
    return {issue.code for issue in exc_info.value.issues}


def test_equivalent_xlsx_and_json_hash_match() -> None:
    json_envelope = parse_upload(json_template(), "request.json")
    xlsx_envelope = parse_upload(workbook_template("starter-xlsx"), "request.xlsx")

    assert json_envelope.request["request_id"] == xlsx_envelope.request["request_id"]
    assert json_envelope.summary.row_counts == xlsx_envelope.summary.row_counts
    assert json_envelope.request_hash == xlsx_envelope.request_hash


@pytest.mark.parametrize(
    ("data", "filename", "expected_code"),
    [
        (b"", "request.json", "EMPTY_UPLOAD"),
        (b"not json or xlsx", "request.csv", "UPLOAD_TYPE"),
        (b"\xff\xfe not utf8", "request.json", "JSON_ENCODING"),
        (b"{not valid json", "request.json", "JSON_SYNTAX"),
        (b"[]", "request.json", "JSON_ROOT_TYPE"),
        (b"garbage-not-a-zip", "request.xlsx", "XLSX_ARCHIVE_INVALID"),
    ],
)
def test_upload_rejection_matrix(data: bytes, filename: str, expected_code: str) -> None:
    with pytest.raises(UploadValidationError) as exc_info:
        parse_upload(data, filename)

    assert expected_code in _issue_codes(exc_info)


def test_upload_size_limit_is_enforced() -> None:
    oversized = b"0" * (10 * 1024 * 1024 + 1)

    with pytest.raises(UploadValidationError) as exc_info:
        parse_upload(oversized, "request.json")

    assert _issue_codes(exc_info) == {"UPLOAD_SIZE"}


def test_json_upload_reports_contract_issues_for_missing_required_field() -> None:
    payload = json.loads(json_template())
    del payload["request_id"]

    with pytest.raises(UploadValidationError) as exc_info:
        parse_upload(json.dumps(payload).encode("utf-8"), "request.json")

    assert any(issue.code.startswith("CONTRACT_") for issue in exc_info.value.issues)


def test_json_upload_rejects_duplicate_keys() -> None:
    with pytest.raises(UploadValidationError) as exc_info:
        parse_upload(b'{"request_id":"one","request_id":"two"}', "bad.json")

    assert _issue_codes(exc_info) == {"DUPLICATE_JSON_KEY"}


def test_xlsx_upload_rejects_missing_and_unknown_sheets() -> None:
    workbook = load_workbook(BytesIO(workbook_template("starter-xlsx")))
    workbook.remove(workbook["demand"])
    workbook.create_sheet("mystery_sheet")
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(UploadValidationError) as exc_info:
        parse_upload(output.getvalue(), "bad.xlsx")

    codes = _issue_codes(exc_info)
    assert "MISSING_SHEET" in codes
    assert "UNKNOWN_SHEET" in codes


def test_xlsx_upload_rejects_unknown_and_missing_and_duplicate_columns() -> None:
    workbook = load_workbook(BytesIO(workbook_template("starter-xlsx")))
    sheet = workbook["inventory"]
    sheet["A1"] = "not_a_real_column"
    sheet["B1"] = sheet["C1"].value  # duplicate an existing header

    with pytest.raises(UploadValidationError) as exc_info:
        parse_upload(_save(workbook), "bad.xlsx")

    codes = _issue_codes(exc_info)
    assert "UNKNOWN_COLUMN" in codes
    assert "MISSING_COLUMN" in codes
    assert "DUPLICATE_COLUMN" in codes


def test_xlsx_upload_rejects_required_cell_empty_and_bad_cell_type() -> None:
    workbook = load_workbook(BytesIO(workbook_template("starter-xlsx")))
    sheet = workbook["inventory"]
    sheet["F2"] = None  # total_lendable_shares, required number
    sheet["M2"] = "not-a-boolean"  # eligible, boolean

    with pytest.raises(UploadValidationError) as exc_info:
        parse_upload(_save(workbook), "bad.xlsx")

    codes = _issue_codes(exc_info)
    assert "REQUIRED_CELL_EMPTY" in codes
    assert "CELL_TYPE" in codes


def test_xlsx_upload_rejects_formula_in_request_sheet() -> None:
    workbook = load_workbook(BytesIO(workbook_template("starter-xlsx")), data_only=False)
    workbook["request"]["B2"] = '=CONCAT("SYN", "-REQ")'

    with pytest.raises(UploadValidationError) as exc_info:
        parse_upload(_save(workbook), "bad.xlsx")

    assert "FORMULA_NOT_ALLOWED" in _issue_codes(exc_info)


def test_xlsx_upload_rejects_row_limit_overflow() -> None:
    workbook = load_workbook(BytesIO(workbook_template("advanced-xlsx")))
    sheet = workbook["inventory"]
    for offset in range(501):
        sheet.cell(row=2 + offset, column=1, value=f"SYN-INV-OVERFLOW-{offset:04d}")

    with pytest.raises(UploadValidationError) as exc_info:
        parse_upload(_save(workbook), "bad.xlsx")

    assert "ROW_LIMIT" in _issue_codes(exc_info)


def test_xlsx_upload_rejects_unsafe_archive_content() -> None:
    clean = workbook_template("starter-xlsx")
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(clean)) as source, zipfile.ZipFile(
        output, mode="w", compression=zipfile.ZIP_DEFLATED
    ) as sink:
        for item in source.infolist():
            sink.writestr(item, source.read(item.filename))
        sink.writestr("xl/vbaProject.bin", b"not-really-a-macro")

    with pytest.raises(UploadValidationError) as exc_info:
        parse_upload(output.getvalue(), "bad.xlsx")

    assert "XLSX_UNSAFE_CONTENT" in _issue_codes(exc_info)


def test_upload_aggregates_located_independent_errors_without_repair_or_row_drop() -> None:
    workbook = load_workbook(BytesIO(workbook_template("starter-xlsx")), data_only=False)
    workbook["routes"]["A2"] = '=CONCAT("SYN", "-RT")'
    workbook["demand"]["A1"] = "not_a_real_column"

    with pytest.raises(UploadValidationError) as exc_info:
        parse_upload(_save(workbook), "bad.xlsx")

    codes = _issue_codes(exc_info)
    assert "FORMULA_NOT_ALLOWED" in codes
    assert "UNKNOWN_COLUMN" in codes
    assert "MISSING_COLUMN" in codes
    assert len(exc_info.value.issues) >= 3


def _save(workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
